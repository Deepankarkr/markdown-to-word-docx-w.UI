import hashlib
import html
import io
import random
import re
import zipfile
from pathlib import Path

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from flask import Flask, render_template, request, send_file
from markdown import markdown
from markupsafe import Markup
from pygments.formatters import HtmlFormatter
from pypdf import PdfReader, PdfWriter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.graphics.shapes import Circle, Drawing, String
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 32 * 1024 * 1024
ALLOWED_EXTENSIONS = {"md", "markdown", "txt"}
PDF_EXTENSION = "pdf"
INK, PAPER = "263238", "FFFDF3"
ACCENTS = ["E85D75", "3A86FF", "2A9D8F", "F4A261", "8E6CCF", "E76F51"]
PASTELS = ["FFE1E8", "DDEEFF", "DDF4EC", "FFF0CC", "ECE2FF", "FFE5D5"]
HIGHLIGHTS = [WD_COLOR_INDEX.YELLOW, WD_COLOR_INDEX.BRIGHT_GREEN, WD_COLOR_INDEX.TURQUOISE]
BULLET_MARKERS = ["★", "➜", "◆", "✓", "💡", "📌"]
LIST_LINE = re.compile(r"^\s{0,3}(?:[-+*]|\d+[.)])\s+\S")
FENCE_LINE = re.compile(r"^\s{0,3}(`{3,}|~{3,})")


def preserve_mermaid_fences(source):
    """Keep Mermaid's language marker intact while Pygments handles other fences."""
    fence = re.compile(
        r"(?ms)^(?P<indent>[ \t]*)(?P<marker>`{3,}|~{3,})[ \t]*mermaid[^\n]*\n"
        r"(?P<body>.*?)^(?P=indent)(?P=marker)[ \t]*$"
    )

    def replace(match):
        diagram = html.escape(match.group("body").rstrip())
        return f'\n<pre><code class="language-mermaid">{diagram}</code></pre>\n'

    return fence.sub(replace, source)


def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


def clean_name(filename):
    # Keep the user's filename intact while discarding any client-supplied path.
    leaf_name = (filename or "document.md").replace("\\", "/").rsplit("/", 1)[-1]
    leaf_name = "".join(char for char in leaf_name if ord(char) >= 32).strip()
    return Path(leaf_name or "document.md").stem or "document"


def is_pdf(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() == PDF_EXTENSION


def merge_pdf_files(files):
    """Merge PDF page streams without rendering or recompressing their contents."""
    writer = PdfWriter()
    for file in files:
        reader = PdfReader(file.stream, strict=False)
        if reader.is_encrypted:
            raise ValueError(f'"{file.filename}" is password protected.')
        writer.append(reader)

    output = io.BytesIO()
    writer.write(output)
    writer.close()
    output.seek(0)
    return output


def seeded_rng(text):
    seed = int.from_bytes(hashlib.sha256(text.encode("utf-8")).digest()[:8], "big")
    return random.Random(seed)


def normalize_compact_lists(md_text):
    """Add parser-only separation before compact lists without changing content.

    Common Markdown frequently places a list directly below a sentence or bold
    label. Python-Markdown requires a blank line there, otherwise it emits the
    markers as paragraph text. The inserted blank line is structural only and
    never appears as document content. Fenced code is deliberately untouched.
    """
    lines = md_text.splitlines(keepends=True)
    output = []
    active_fence = None
    previous_text = ""
    for line in lines:
        plain = line.rstrip("\r\n")
        fence = FENCE_LINE.match(plain)
        if fence:
            token = fence.group(1)
            if active_fence is None:
                active_fence = token[0]
            elif token[0] == active_fence:
                active_fence = None

        if (active_fence is None and LIST_LINE.match(plain) and previous_text.strip()
                and not LIST_LINE.match(previous_text)
                and not FENCE_LINE.match(previous_text)):
            output.append("\n")
        output.append(line)
        previous_text = plain
    return "".join(output)


def register_pdf_fonts():
    """Use MV Boli when installed, with a portable fallback for deployment."""
    font_path = Path(r"C:\Windows\Fonts\mvboli.ttf")
    if font_path.exists() and "MV Boli" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("MV Boli", str(font_path)))
    return "MV Boli" if "MV Boli" in pdfmetrics.getRegisteredFontNames() else "Helvetica"


def pdf_styles():
    font = register_pdf_fonts()
    styles = getSampleStyleSheet()
    return {
        "font": font,
        "body": ParagraphStyle("NotesBody", parent=styles["BodyText"], fontName=font, fontSize=12,
                               leading=16, textColor=colors.HexColor(f"#{INK}"), spaceAfter=9),
        "h1": ParagraphStyle("NotesH1", parent=styles["Heading1"], fontName=font, fontSize=23,
                             leading=25, textColor=colors.HexColor(f"#{ACCENTS[1]}"), spaceBefore=12, spaceAfter=9),
        "h2": ParagraphStyle("NotesH2", parent=styles["Heading2"], fontName=font, fontSize=18,
                             leading=20, textColor=colors.HexColor(f"#{ACCENTS[2]}"), spaceBefore=10, spaceAfter=7),
        "h3": ParagraphStyle("NotesH3", parent=styles["Heading3"], fontName=font, fontSize=15,
                             leading=17, textColor=colors.HexColor(f"#{ACCENTS[4]}"), spaceBefore=8, spaceAfter=6),
        "code": ParagraphStyle("NotesCode", parent=styles["Code"], fontName="Courier", fontSize=9,
                               leading=12, leftIndent=12, rightIndent=8, spaceBefore=6, spaceAfter=9,
                               backColor=colors.HexColor("#F0EEF6"), borderColor=colors.HexColor("#8E6CCF"),
                               borderWidth=1, borderPadding=7),
        "quote": ParagraphStyle("NotesQuote", parent=styles["BodyText"], fontName=font, fontSize=12,
                                leading=16, leftIndent=16, textColor=colors.HexColor("#5D4777"),
                                backColor=colors.HexColor("#ECE2FF"), borderColor=colors.HexColor("#8E6CCF"),
                                borderWidth=1, borderPadding=7, spaceAfter=9),
        "bullet": ParagraphStyle("NotesBullet", parent=styles["BodyText"], fontName=font, fontSize=12,
                                 leading=16, textColor=colors.HexColor(f"#{INK}")),
        "marker": ParagraphStyle("NotesMarker", parent=styles["BodyText"], fontName=font, fontSize=16,
                                 leading=16, alignment=TA_CENTER),
    }


def pdf_inline(element):
    parts = []
    for child in element.children:
        if isinstance(child, NavigableString):
            parts.append(html.escape(str(child)))
        elif isinstance(child, Tag):
            if child.name in ("ul", "ol"):
                continue
            if child.name == "br":
                parts.append("<br/>")
                continue
            content = pdf_inline(child)
            if child.name in ("strong", "b"):
                content = f"<b>{content}</b>"
            elif child.name in ("em", "i"):
                content = f"<i><u>{content}</u></i>"
            elif child.name == "code":
                content = f'<font name="Courier" color="#9C2F56">{content}</font>'
            elif child.name == "a":
                href = html.escape(child.get("href", ""), quote=True)
                content = f'<link href="{href}" color="#3A86FF"><u>{content}</u></link>'
            parts.append(content)
    return "".join(parts)


def add_pdf_list(story, list_tag, styles, rng, ordered=False, depth=0):
    for index, item in enumerate(list_tag.find_all("li", recursive=False), start=1):
        content = Paragraph(pdf_inline(item), styles["bullet"])
        if ordered:
            marker = Paragraph(f"{index}.", styles["marker"])
            fill, accent = colors.white, colors.HexColor(f"#{ACCENTS[1]}")
        else:
            palette = rng.randrange(len(ACCENTS))
            fill, accent = colors.HexColor(f"#{PASTELS[palette]}"), colors.HexColor(f"#{ACCENTS[palette]}")
            # Draw the marker as vector artwork. This remains crisp and cannot
            # turn into a missing-glyph square in different PDF viewers.
            marker_text = ">"
            marker = Drawing(24, 24)
            marker.add(Circle(12, 12, 10, fillColor=accent, strokeColor=None))
            marker.add(String(12, 7.5, marker_text, fontName="Helvetica-Bold", fontSize=13,
                              fillColor=colors.white, textAnchor="middle"))
        width = max(3.8 * inch, 6.55 * inch - depth * 0.28 * inch)
        card = Table([[marker, content]], colWidths=[0.42 * inch, width - 0.42 * inch], hAlign="LEFT")
        card.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), fill), ("LINEBEFORE", (0, 0), (0, -1), 4, accent),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 7),
            ("RIGHTPADDING", (0, 0), (-1, -1), 7), ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        story.extend([card, Spacer(1, 5)])
        for nested in item.find_all(["ul", "ol"], recursive=False):
            add_pdf_list(story, nested, styles, rng, nested.name == "ol", depth + 1)


def add_markdown_to_pdf(story, md_text, styles, source_title=None):
    rng = seeded_rng((source_title or "") + "\n" + md_text)
    soup = BeautifulSoup(markdown(normalize_compact_lists(md_text),
        extensions=["tables", "fenced_code", "sane_lists", "nl2br"], output_format="html5"), "html.parser")
    if source_title:
        story.append(Paragraph(html.escape(source_title), styles["h1"]))
    for element in soup.children:
        if not isinstance(element, Tag):
            continue
        if re.fullmatch(r"h[1-6]", element.name):
            story.append(Paragraph(pdf_inline(element), styles[f"h{min(int(element.name[1]), 3)}"]))
        elif element.name == "p":
            story.append(Paragraph(pdf_inline(element), styles["body"]))
        elif element.name in ("ul", "ol"):
            add_pdf_list(story, element, styles, rng, element.name == "ol")
        elif element.name == "pre":
            story.append(Paragraph(html.escape(element.get_text()).replace("\n", "<br/>"), styles["code"]))
        elif element.name == "blockquote":
            story.append(Paragraph(html.escape(element.get_text(" ", strip=True)), styles["quote"]))
        elif element.name == "table":
            rows = element.find_all("tr")
            data = [[Paragraph(html.escape(cell.get_text(" ", strip=True)), styles["bullet"])
                     for cell in row.find_all(["th", "td"])] for row in rows]
            if data:
                table = Table(data, repeatRows=1, hAlign="LEFT")
                commands = [("GRID", (0, 0), (-1, -1), 0.8, colors.HexColor("#8E6CCF")),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#3A86FF")),
                            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("TOPPADDING", (0, 0), (-1, -1), 7),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 7)]
                for row_index in range(1, len(data)):
                    commands.append(("BACKGROUND", (0, row_index), (-1, row_index), colors.HexColor(f"#{PASTELS[row_index % len(PASTELS)]}")))
                table.setStyle(TableStyle(commands))
                story.extend([table, Spacer(1, 9)])
        elif element.name == "hr":
            story.append(Paragraph("★ &nbsp; ~ &nbsp; ★ &nbsp; ~ &nbsp; ★", styles["h3"]))


def pdf_page(canvas, doc):
    canvas.saveState()
    canvas.setFillColor(colors.HexColor(f"#{PAPER}"))
    canvas.rect(0, 0, A4[0], A4[1], fill=1, stroke=0)
    canvas.restoreState()


def make_pdf(md_text, title):
    output, story, styles = io.BytesIO(), [], pdf_styles()
    title_text = Paragraph(html.escape(title), styles["h1"])
    title_card = Table([[title_text]], colWidths=[6.55 * inch], hAlign="CENTER")
    title_card.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(f"#{PASTELS[1]}")),
        ("BOX", (0, 0), (-1, -1), 2, colors.HexColor(f"#{ACCENTS[0]}")),
        ("TOPPADDING", (0, 0), (-1, -1), 10), ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
        ("LEFTPADDING", (0, 0), (-1, -1), 12), ("RIGHTPADDING", (0, 0), (-1, -1), 12),
    ]))
    story.extend([title_card, Spacer(1, 14)])
    add_markdown_to_pdf(story, md_text, styles)
    document = SimpleDocTemplate(output, pagesize=A4, title=title, topMargin=0.55 * inch,
                                 bottomMargin=0.55 * inch, leftMargin=0.65 * inch, rightMargin=0.65 * inch)
    document.build(story, onFirstPage=pdf_page, onLaterPages=pdf_page)
    output.seek(0)
    return output


def make_merged_pdf(files):
    output, story, styles = io.BytesIO(), [], pdf_styles()
    for index, file in enumerate(files):
        if index:
            story.append(PageBreak())
        add_markdown_to_pdf(story, file.read().decode("utf-8", errors="replace"), styles, clean_name(file.filename))
    document = SimpleDocTemplate(output, pagesize=A4, title="Merged Markdown Document", topMargin=0.55 * inch,
                                 bottomMargin=0.55 * inch, leftMargin=0.65 * inch, rightMargin=0.65 * inch)
    document.build(story, onFirstPage=pdf_page, onLaterPages=pdf_page)
    output.seek(0)
    return output


def set_font(run, name="MV Boli", size=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_page_color(doc, color=PAPER):
    background = OxmlElement("w:background")
    background.set(qn("w:color"), color)
    doc._element.insert(0, background)
    doc.settings._element.append(OxmlElement("w:displayBackgroundShape"))


def set_paragraph_border(paragraph, color, sides=("bottom",), size="10", style="single"):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    for side in sides:
        edge = OxmlElement(f"w:{side}")
        for key, value in (("val", style), ("sz", size), ("space", "5"), ("color", color)):
            edge.set(qn(f"w:{key}"), value)
        borders.append(edge)


def shade_paragraph(paragraph, fill):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    paragraph._p.get_or_add_pPr().append(shading)


def shade_cell(cell, fill):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_margins(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", 110), ("start", 130), ("bottom", 110), ("end", 130)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)


def setup_document(doc):
    section = doc.sections[0]
    section.top_margin, section.bottom_margin = Inches(0.72), Inches(0.72)
    section.left_margin, section.right_margin = Inches(0.82), Inches(0.72)
    set_page_color(doc)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name, normal.font.size = "MV Boli", Pt(13)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "MV Boli")
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after, normal.paragraph_format.line_spacing = Pt(10), 1.0
    for level, size in ((1, 25), (2, 20), (3, 17)):
        style = styles[f"Heading {level}"]
        style.font.name, style.font.size, style.font.bold = "MV Boli", Pt(size), True
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "MV Boli")
        style.font.color.rgb = RGBColor.from_string(ACCENTS[level - 1])
        style.paragraph_format.space_before = Pt(18 if level == 1 else 13)
        style.paragraph_format.space_after, style.paragraph_format.keep_with_next = Pt(9), True
    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name, style.font.size = "MV Boli", Pt(13)
        style.paragraph_format.space_after, style.paragraph_format.line_spacing = Pt(7), 1.0
    if "Handwritten Bullet" not in [s.name for s in styles]:
        style = styles.add_style("Handwritten Bullet", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
        style.font.name, style.font.size, style.font.bold = "MV Boli", Pt(13.5), True
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "MV Boli")
        style.paragraph_format.space_before = Pt(4)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_together = True
    if "Code Block" not in [s.name for s in styles]:
        style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name, style.font.size = "Consolas", Pt(10.5)
        style.font.color.rgb = RGBColor.from_string("39424E")
        style.paragraph_format.left_indent, style.paragraph_format.right_indent = Inches(0.22), Inches(0.16)
        style.paragraph_format.space_before, style.paragraph_format.space_after = Pt(8), Pt(12)
        style.paragraph_format.line_spacing = 1.0
    if "Quote Block" not in [s.name for s in styles]:
        style = styles.add_style("Quote Block", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name, style.font.size, style.font.italic = "MV Boli", Pt(13), True
        style.font.color.rgb = RGBColor.from_string("5D4777")
        style.paragraph_format.left_indent, style.paragraph_format.right_indent = Inches(0.28), Inches(0.12)
        style.paragraph_format.space_before, style.paragraph_format.space_after = Pt(8), Pt(12)
        style.paragraph_format.line_spacing = 1.0


def add_title_page(doc, title, subtitle=None):
    doodle = doc.add_paragraph()
    doodle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(doodle.add_run("✦  ✎  ✦"), size=17, color=ACCENTS[3])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(8), Pt(15)
    run = p.add_run(title)
    run.bold = True
    set_font(run, size=29, color=ACCENTS[1])
    shade_paragraph(p, PASTELS[1])
    set_paragraph_border(p, ACCENTS[0], ("top", "bottom"), "14", "double")
    if subtitle:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(20)
        set_font(p.add_run(subtitle), size=12, color="70647A")


def add_inline_text(paragraph, element, rng):
    """Copy inline text exactly once and apply presentation-only emphasis."""
    for child in element.children:
        if isinstance(child, NavigableString):
            paragraph.add_run(str(child))
        elif isinstance(child, Tag):
            if child.name == "br":
                paragraph.add_run().add_break()
                continue
            # Nested lists are emitted separately by add_list, never folded into
            # their parent item (which would duplicate their text).
            if child.name in ("ul", "ol"):
                continue
            run = paragraph.add_run(child.get_text())
            if child.name in ("strong", "b"):
                run.bold = True
                run.font.highlight_color = rng.choice(HIGHLIGHTS)
            if child.name in ("em", "i"):
                run.italic, run.underline = True, True
                run.font.color.rgb = RGBColor.from_string(rng.choice(ACCENTS))
            if child.name == "code":
                set_font(run, "Consolas", 10.5, "9C2F56")
                run.font.highlight_color = WD_COLOR_INDEX.GRAY_25
            if child.name == "a":
                run.underline = True
                run.font.color.rgb = RGBColor.from_string(ACCENTS[1])


def style_heading(paragraph, level, rng):
    index = rng.randrange(len(ACCENTS))
    for run in paragraph.runs:
        set_font(run, size={1: 25, 2: 20, 3: 17}[level], color=ACCENTS[index])
        run.bold = True
    if level == 1:
        shade_paragraph(paragraph, PASTELS[index])
        set_paragraph_border(paragraph, ACCENTS[index], ("left", "bottom"), "13")
    elif level == 2:
        set_paragraph_border(paragraph, ACCENTS[index], ("bottom",), "9", "dashed")
    else:
        set_paragraph_border(paragraph, ACCENTS[index], ("left",), "12")


def set_cell_text(cell, text, bold, color, fill):
    cell.text = ""
    shade_cell(cell, fill)
    set_cell_margins(cell)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after, p.paragraph_format.line_spacing = Pt(2), 1.0
    run = p.add_run(text)
    run.bold = bold
    set_font(run, size=12.5 if bold else 12, color=color)


def add_table(doc, table_tag, rng):
    rows = table_tag.find_all("tr")
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=max(len(row.find_all(["th", "td"])) for row in rows))
    table.style, table.autofit = "Table Grid", True
    offset = rng.randrange(len(PASTELS))
    for row_index, row in enumerate(rows):
        for col_index, cell in enumerate(row.find_all(["th", "td"])):
            header = cell.name == "th" or (row_index == 0 and row.find("th") is not None)
            index = (offset + col_index + row_index // 2) % len(PASTELS)
            set_cell_text(table.cell(row_index, col_index), cell.get_text(" ", strip=True), header,
                          "FFFFFF" if header else INK, ACCENTS[index] if header else PASTELS[index])
    doc.add_paragraph()


def add_list(doc, list_tag, rng, ordered=False, depth=0):
    for item in list_tag.find_all("li", recursive=False):
        p = doc.add_paragraph(style="List Number" if ordered else "Handwritten Bullet")
        if not ordered:
            accent_index = rng.randrange(len(ACCENTS))
            p.paragraph_format.left_indent = Inches(0.42 + depth * 0.3)
            p.paragraph_format.right_indent = Inches(0.08)
            p.paragraph_format.first_line_indent = Inches(-0.3)
            marker_text = rng.choice(BULLET_MARKERS)
            marker = p.add_run(marker_text + "  ")
            marker.bold = True
            marker_font = "Segoe UI Emoji" if marker_text in ("💡", "📌") else "Segoe UI Symbol"
            set_font(marker, marker_font, 18, ACCENTS[accent_index])
            shade_paragraph(p, PASTELS[accent_index])
            set_paragraph_border(p, ACCENTS[accent_index], ("left",), "16")
        add_inline_text(p, item, rng)
        for run in p.runs:
            if (ordered or run is not marker) and run.font.color.rgb is None:
                run.font.color.rgb = RGBColor.from_string(INK if not ordered else rng.choice(ACCENTS))
        for nested in item.find_all(["ul", "ol"], recursive=False):
            add_list(doc, nested, rng, nested.name == "ol", depth + 1)


def add_markdown_to_doc(doc, md_text, source_title=None):
    rng = seeded_rng((source_title or "") + "\n" + md_text)
    parse_text = normalize_compact_lists(md_text)
    html = markdown(parse_text, extensions=["tables", "fenced_code", "sane_lists", "nl2br"], output_format="html5")
    soup = BeautifulSoup(html, "html.parser")
    if source_title:
        heading = doc.add_heading(source_title, level=1)
        style_heading(heading, 1, rng)
    for element in soup.children:
        if not isinstance(element, Tag):
            continue
        if re.fullmatch(r"h[1-6]", element.name):
            level = min(int(element.name[1]), 3)
            p = doc.add_heading(level=level)
            add_inline_text(p, element, rng)
            style_heading(p, level, rng)
        elif element.name == "p":
            add_inline_text(doc.add_paragraph(), element, rng)
        elif element.name in ("ul", "ol"):
            add_list(doc, element, rng, element.name == "ol")
        elif element.name == "pre":
            p = doc.add_paragraph(style="Code Block")
            p.add_run(element.get_text())
            shade_paragraph(p, rng.choice(PASTELS))
            set_paragraph_border(p, rng.choice(ACCENTS), ("top", "left", "bottom", "right"), "7", "dashed")
        elif element.name == "blockquote":
            p = doc.add_paragraph(style="Quote Block")
            p.add_run(element.get_text(" ", strip=True))
            shade_paragraph(p, PASTELS[4])
            set_paragraph_border(p, ACCENTS[4], ("left", "bottom"), "12")
        elif element.name == "table":
            add_table(doc, element, rng)
        elif element.name == "hr":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(p.add_run("✦  ~  ✦  ~  ✦"), size=14, color=rng.choice(ACCENTS))


def make_docx(md_text, title, include_title=True):
    doc = Document()
    setup_document(doc)
    if include_title:
        add_title_page(doc, title)
    add_markdown_to_doc(doc, md_text)
    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output


def make_merged_docx(files):
    doc = Document()
    setup_document(doc)
    for index, file in enumerate(files):
        filename = file.filename
        text = file.read().decode("utf-8", errors="replace")
        if index:
            doc.add_section(WD_SECTION.NEW_PAGE)
        add_markdown_to_doc(doc, text, clean_name(filename))
    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/pdf")
def pdf_tools():
    return render_template("pdf.html")


@app.post("/merge-pdf")
def merge_pdfs():
    files = request.files.getlist("pdf_files")
    valid = [file for file in files if file and file.filename and is_pdf(file.filename)]
    if len(valid) < 2:
        return render_template("pdf.html", error="Please select at least two PDF files."), 400

    try:
        output = merge_pdf_files(valid)
    except (ValueError, OSError, EOFError) as exc:
        return render_template("pdf.html", error=f"Could not merge the PDFs: {exc}"), 400
    except Exception:
        return render_template(
            "pdf.html", error="Could not merge the PDFs. Check that every file is a valid, unprotected PDF."
        ), 400

    return send_file(
        output,
        as_attachment=True,
        download_name="merged_document.pdf",
        mimetype="application/pdf",
    )


@app.post("/view")
def view_markdown():
    files = request.files.getlist("files")
    valid = [file for file in files if file and file.filename and allowed_file(file.filename)]
    if len(valid) != 1:
        return "Select exactly one Markdown file to preview.", 400

    file = valid[0]
    source = file.read().decode("utf-8", errors="replace")
    rendered = markdown(
        preserve_mermaid_fences(source),
        extensions=["fenced_code", "tables", "toc", "sane_lists", "smarty", "codehilite"],
        extension_configs={
            "codehilite": {"guess_lang": False, "css_class": "codehilite", "linenums": False},
            "toc": {"permalink": True},
        },
        output_format="html5",
    )
    title_match = re.search(r"^#\s+(.+)$", source, re.MULTILINE)
    title = re.sub(r"[*_`\[\]]", "", title_match.group(1)).strip() if title_match else clean_name(file.filename)
    return render_template(
        "viewer.html",
        title=title,
        filename=file.filename,
        content=Markup(rendered),
        pygments_css=Markup(HtmlFormatter(style="monokai").get_style_defs(".codehilite")),
    )


@app.route("/convert", methods=["POST"])
def convert():
    files = request.files.getlist("files")
    valid = [f for f in files if f and f.filename and allowed_file(f.filename)]
    output_format = request.form.get("output_format", "docx").lower()
    if output_format not in ("docx", "pdf"):
        output_format = "docx"
    if not valid:
        return render_template("index.html", error="Please upload at least one .md, .markdown, or .txt file."), 400
    if len(valid) == 1:
        file = valid[0]
        base = clean_name(file.filename)
        text = file.read().decode("utf-8", errors="replace")
        if output_format == "pdf":
            return send_file(make_pdf(text, base), as_attachment=True, download_name=f"{base}.pdf",
                             mimetype="application/pdf")
        return send_file(
            make_docx(text, base),
            as_attachment=True,
            download_name=f"{base}.docx",
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    if request.form.get("mode", "separate") == "merged":
        if output_format == "pdf":
            return send_file(make_merged_pdf(valid), as_attachment=True, download_name="merged_markdown_document.pdf",
                             mimetype="application/pdf")
        return send_file(make_merged_docx(valid), as_attachment=True, download_name="merged_markdown_document.docx",
                         mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for file in valid:
            text, base = file.read().decode("utf-8", errors="replace"), clean_name(file.filename)
            if output_format == "pdf":
                archive.writestr(f"{base}.pdf", make_pdf(text, base).getvalue())
            else:
                archive.writestr(f"{base}.docx", make_docx(text, base).getvalue())
    output.seek(0)
    zip_name = "converted_pdf_files.zip" if output_format == "pdf" else "converted_docx_files.zip"
    return send_file(output, as_attachment=True, download_name=zip_name, mimetype="application/zip")


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5000, debug=True)
